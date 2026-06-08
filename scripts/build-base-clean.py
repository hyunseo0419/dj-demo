#!/usr/bin/env python3
"""
첨부 이미지 형태의 섹션 셀을 '처음부터 깨끗하게' 생성한다(손상된 원본 미사용).
python-docx로 무결한 최소 docx를 만들고, 거기서 base.docx + section-unit.xml을 뽑는다.

섹션 셀 구조(이미지와 동일):
  ┌────┬───────────────────────────┐
  │번호 │ 관찰 텍스트(번호 옆)  {{DESC}} │  (col1+col2 병합)
  │    ├──────────────┬────────────┤
  │    │ Before        │ After       │
  │    │ [이미지]       │ (빈칸)       │
  │    │ Date of Rectified (yyyy/mm/dd): │
  │    │ (빈 행)                        │  (col1+col2 병합)
  └────┴───────────────────────────┘
좌측 번호열(col0)은 세로 병합. 출력물의 letterhead/상단행은 추후 완벽 양식에서.
"""
import re, os, zipfile, shutil, tempfile
from docx import Document
from docx.shared import Cm, Emu
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "public/template")
PLACEHOLDER_IMG = "/tmp/testimgs/image1.jpeg"  # extent 산출용(런타임에 교체됨)

# 원본 양식과 동일한 컬럼 폭(dxa): 좁은 번호칸 + 동일한 Before/After 절반
DXA = [421, 4606, 4607]
TOTAL = sum(DXA)            # 9634
DXA2EMU = 635               # 1 dxa = 635 EMU
NUM_W = Emu(DXA[0] * DXA2EMU)
HALF1 = Emu(DXA[1] * DXA2EMU)
HALF2 = Emu(DXA[2] * DXA2EMU)

doc = Document()
# 페이지 여백만 살짝(기본도 무방)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2)

table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"           # 전체 테두리(깨끗한 기본 스타일)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

# 병합 먼저: 번호열 세로 병합 + R0/R4 가로 병합
num_cell = table.cell(0, 0)
for r in range(1, 5):
    num_cell = num_cell.merge(table.cell(r, 0))
obs_cell = table.cell(0, 1).merge(table.cell(0, 2))
bottom_cell = table.cell(4, 1).merge(table.cell(4, 2))

# 폭 지정: 고정 레이아웃 + 셀/그리드 폭(dxa)
num_cell.width = NUM_W
obs_cell.width = Emu(TOTAL * DXA2EMU - NUM_W)  # col1+col2 span
bottom_cell.width = Emu(TOTAL * DXA2EMU - NUM_W)
for r in (1, 2, 3):
    table.cell(r, 1).width = HALF1
    table.cell(r, 2).width = HALF2

# 고정 레이아웃 + 전체 폭 + tblGrid 컬럼폭(좁은 번호칸이 유지되도록)
tbl = table._tbl
tblPr = tbl.tblPr
for tag, attrs in (
    ("w:tblW", {"w:type": "dxa", "w:w": str(TOTAL)}),
    ("w:tblLayout", {"w:type": "fixed"}),
):
    el = tblPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        tblPr.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
grid = tbl.find(qn("w:tblGrid"))
for col, w in zip(grid.findall(qn("w:gridCol")), DXA):
    col.set(qn("w:w"), str(w))

# 테두리를 '스타일'이 아니라 표에 '직접' 박는다(어떤 뷰어/워드에서도 선이 보이도록)
borders = OxmlElement("w:tblBorders")
for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
    e = OxmlElement("w:" + edge)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), "8")        # 약 1pt
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), "000000")
    borders.append(e)
old = tblPr.find(qn("w:tblBorders"))
if old is not None:
    tblPr.remove(old)
tblPr.append(borders)

# 내용
num_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
num_cell.paragraphs[0].text = "1"                      # -> {{NUMBER}}
obs_cell.paragraphs[0].text = "OBS_PLACEHOLDER_TEXT"   # -> {{DESC}}

for col, label in ((1, "Before"), (2, "After")):
    c = table.cell(1, col)
    c.paragraphs[0].text = label
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 이미지(Before 셀). 폭 8cm로 넣고 extent는 뒤에서 placeholder로 치환
img_p = table.cell(2, 1).paragraphs[0]
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.add_run().add_picture(PLACEHOLDER_IMG, width=Cm(8))

table.cell(3, 1).paragraphs[0].text = "Date of Rectified (yyyy/mm/dd):"

tmp = tempfile.mkdtemp()
clean_docx = os.path.join(tmp, "clean.docx")
doc.save(clean_docx)

# ---- 추출 + 파라미터화 ----
work = os.path.join(tmp, "x")
with zipfile.ZipFile(clean_docx) as z:
    z.extractall(work)
doc_path = os.path.join(work, "word/document.xml")
x = open(doc_path, encoding="utf-8").read()

m = re.search(r"<w:tbl>.*?</w:tbl>", x, re.S)
unit = m.group(0)

# 번호/설명 텍스트 치환
unit = unit.replace("<w:t>1</w:t>", "<w:t>{{NUMBER}}</w:t>", 1)
unit = re.sub(r"(<w:t[^>]*>)OBS_PLACEHOLDER_TEXT(</w:t>)", r"\1{{DESC}}\2", unit, count=1)
# 이미지 drawing 파라미터화
embed = re.search(r'r:embed="(rId\d+)"', unit).group(1)
ext = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', unit)
cx, cy = ext.group(1), ext.group(2)
docpr = re.search(r'<wp:docPr id="(\d+)"', unit)
unit = unit.replace('r:embed="%s"' % embed, 'r:embed="{{RID}}"')
unit = unit.replace('cx="%s" cy="%s"' % (cx, cy), 'cx="{{CX}}" cy="{{CY}}"')
if docpr:
    unit = unit.replace('id="%s"' % docpr.group(1), 'id="{{DOCPR_ID}}"', 1)

# base 문서 = 표 제거 + 마커
base_doc = x[: m.start()] + "<!--GEN_SECTIONS-->" + x[m.end():]
open(doc_path, "w", encoding="utf-8").write(base_doc)

# 이미지 rels 제거(본문) + placeholder 미디어 제거 (Content_Types의 jpeg default는 유지)
rels_path = os.path.join(work, "word/_rels/document.xml.rels")
rels = open(rels_path, encoding="utf-8").read()
rels = re.sub(r'<Relationship [^>]*relationships/image[^>]*/>', "", rels)
open(rels_path, "w", encoding="utf-8").write(rels)
media_dir = os.path.join(work, "word/media")
if os.path.isdir(media_dir):
    for f in os.listdir(media_dir):
        os.remove(os.path.join(media_dir, f))

os.makedirs(OUT_DIR, exist_ok=True)
open(os.path.join(OUT_DIR, "section-unit.xml"), "w", encoding="utf-8").write(unit)
base_out = os.path.join(OUT_DIR, "base.docx")
if os.path.exists(base_out):
    os.remove(base_out)
with zipfile.ZipFile(base_out, "w", zipfile.ZIP_DEFLATED) as z:
    for root_, _, files in os.walk(work):
        for fn in files:
            full = os.path.join(root_, fn)
            z.write(full, os.path.relpath(full, work))
shutil.rmtree(tmp)

ph = {k: unit.count(k) for k in ["{{NUMBER}}", "{{DESC}}", "{{RID}}", "{{CX}}", "{{CY}}", "{{DOCPR_ID}}"]}
print("clean base.docx:", os.path.getsize(base_out), "bytes | placeholders:", ph)
