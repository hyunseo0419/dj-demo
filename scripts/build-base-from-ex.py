#!/usr/bin/env python3
"""
ex.docx(실제 양식)에서 완성본용 base.docx + section-unit.xml 을 추출한다.
 - base.docx: 레터헤드(FMS 로고/제목/Ship's Name 표/안내문) + "No | Observation/Requisition items"
   헤더행(1회) + 섹션 삽입 마커 + sectPr/footer. 본문 섹션/이미지는 제거.
 - section-unit.xml: table1의 섹션 행(R1~R5)을 파라미터화한 클론 단위(번호/관찰텍스트/이미지, Remark 없음).

사용: python3 scripts/build-base-from-ex.py [ex.docx 경로]
"""
import re, os, sys, zipfile, shutil, tempfile

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "ex.docx")
OUT_DIR = os.path.join(ROOT, "public/template")
IMAGE_REL = 'relationships/image'
MARKER = "<!--GEN_SECTIONS-->"

# 상단 헤더 표: 라벨 -> 플레이스홀더 키 (값 셀을 {{KEY}}로 치환)
HEADER_LABELS = {
    "Date:": "DATE",
    "Ship's Name": "SHIP_NAME",
    "IMO Number": "IMO",
    "Flag": "FLAG",
    "Keel Lay": "KEEL_LAY",
    "Delivery": "DELIVERY",
    "Gross Tonnage": "GROSS_TONNAGE",
    "Deadweight": "DEADWEIGHT",
    "Owner": "OWNER",
    "Manning Company": "MANNING_COMPANY",
    "Crew nationality": "CREW_NATIONALITY",
    "Number of Crew": "NUMBER_OF_CREW",
    "Date of Inspection": "DATE_OF_INSPECTION",
    "Port of Inspection": "PORT_OF_INSPECTION",
    "Inspector": "INSPECTOR",
    "Due date to rectify": "DUE_DATE",
    "Date of Rectified": "DATE_OF_RECTIFIED",
    "Master": "MASTER",
    "Chief officer": "CHIEF_OFFICER",
    "Chief engineer": "CHIEF_ENGINEER",
}


def _norm(s):
    return " ".join(s.split())


def _distinct(row):
    seen, out = [], []
    for c in row.cells:
        if c._tc not in seen:
            seen.append(c._tc)
            out.append(c)
    return out


def _set_ph(cell, ph):
    """값 셀 내용을 단일 런 {{KEY}}로 교체(첫 런 서식 유지)."""
    p = cell.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = ph
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(ph)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def parameterize_header(src_path):
    """헤더 표 값 셀을 {{KEY}}로 치환한 임시 docx 경로 반환."""
    doc = Document(src_path)
    ht = next(
        t for t in doc.tables
        if any("Ship's Name" in c.text for r in t.rows for c in r.cells)
    )
    for row in ht.rows:
        cells = _distinct(row)
        i = 0
        while i < len(cells) - 1:
            key = HEADER_LABELS.get(_norm(cells[i].text))
            if key:
                _set_ph(cells[i + 1], "{{%s}}" % key)
                i += 2
            else:
                i += 1
    out = os.path.join(tempfile.mkdtemp(), "param.docx")
    doc.save(out)
    return out


work = tempfile.mkdtemp()
with zipfile.ZipFile(parameterize_header(SRC)) as z:
    z.extractall(work)
doc_path = os.path.join(work, "word/document.xml")
x = open(doc_path, encoding="utf-8").read()

tbls = list(re.finditer(r"<w:tbl>.*?</w:tbl>", x, re.S))
# 첫 섹션 표 = Before/After 포함하는 첫 표(= No/Observation 헤더 + 섹션1이 합쳐진 표)
i1 = next(i for i, m in enumerate(tbls) if "Before" in m.group(0) and "After" in m.group(0))
t1 = tbls[i1].group(0)
head_open = t1[: t1.find("<w:tr")]          # <w:tbl><w:tblPr>..<w:tblGrid>..
rows = re.findall(r"<w:tr[ >].*?</w:tr>", t1, re.S)
assert len(rows) >= 6, "예상과 다른 행 수: %d" % len(rows)

header_table = head_open + rows[0] + "</w:tbl>"          # No | Observation/Requisition items
unit = head_open + "".join(rows[1:]) + "</w:tbl>"         # 번호+관찰 / Before·After / 이미지 / Date / 빈행

# --- 파라미터화 ---
mnum = re.search(r"<w:t[^>]*>(\d+)</w:t>", unit)
unit = unit[: mnum.start()] + mnum.group(0).replace(mnum.group(1), "{{NUMBER}}") + unit[mnum.end():]
# 관찰 텍스트(가장 긴 런) -> {{DESC}}
runs = [m for m in re.finditer(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", unit) if "{{" not in m.group(1)]
desc = max(runs, key=lambda m: len(m.group(1)))
unit = unit[: desc.start(1)] + "{{DESC}}" + unit[desc.end(1):]
# 이미지 drawing
embed = re.search(r'r:embed="(rId\d+)"', unit).group(1)
ext = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', unit)
cx, cy = ext.group(1), ext.group(2)
docpr = re.search(r'<wp:docPr id="(\d+)"', unit)
unit = unit.replace('r:embed="%s"' % embed, 'r:embed="{{RID}}"')
unit = unit.replace('cx="%s" cy="%s"' % (cx, cy), 'cx="{{CX}}" cy="{{CY}}"')
if docpr:
    unit = unit.replace('id="%s"' % docpr.group(1), 'id="{{DOCPR_ID}}"', 1)

# --- base 문서: 레터헤드 + 헤더행 + 마커 + tail ---
prefix = x[: tbls[i1].start()]               # 레터헤드 + "Click..." 문단
tail = x[tbls[-1].end():]                     # 마지막 표 뒤(빈 문단 + sectPr)
base_doc = prefix + header_table + "<w:p/>" + MARKER + tail
open(doc_path, "w", encoding="utf-8").write(base_doc)

# --- 본문 이미지 rels 제거(헤더/푸터/스타일 유지) ---
rels_path = os.path.join(work, "word/_rels/document.xml.rels")
rels = open(rels_path, encoding="utf-8").read()
rels = re.sub(r"<Relationship [^>]*%s[^>]*/>" % IMAGE_REL, "", rels)
open(rels_path, "w", encoding="utf-8").write(rels)

# --- 헤더/푸터가 참조하는 미디어(로고)만 남기고 본문 미디어 제거 ---
referenced = set()
rels_dir = os.path.join(work, "word/_rels")
for rf in os.listdir(rels_dir):
    if rf.startswith(("header", "footer")):
        t = open(os.path.join(rels_dir, rf), encoding="utf-8").read()
        referenced |= set(re.findall(r'Target="media/([^"]+)"', t))
media_dir = os.path.join(work, "word/media")
removed = 0
for f in os.listdir(media_dir):
    if f not in referenced:
        os.remove(os.path.join(media_dir, f)); removed += 1

os.makedirs(OUT_DIR, exist_ok=True)
open(os.path.join(OUT_DIR, "section-unit.xml"), "w", encoding="utf-8").write(unit)
base_out = os.path.join(OUT_DIR, "base.docx")
if os.path.exists(base_out):
    os.remove(base_out)
with zipfile.ZipFile(base_out, "w", zipfile.ZIP_DEFLATED) as z:
    for r_, _, files in os.walk(work):
        for fn in files:
            full = os.path.join(r_, fn)
            z.write(full, os.path.relpath(full, work))
shutil.rmtree(work)

ph = {k: unit.count(k) for k in ["{{NUMBER}}", "{{DESC}}", "{{RID}}", "{{CX}}", "{{CY}}", "{{DOCPR_ID}}"]}
print("base.docx:", os.path.getsize(base_out), "bytes | media kept:", referenced or "none", "removed:", removed)
print("placeholders:", ph, "| Remark in unit:", "Remark" in unit, "| marker:", MARKER in base_doc)
